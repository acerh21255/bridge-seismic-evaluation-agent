"""Render a finalized ReportDoc into the lab master template (bridge_report/report/template)."""
import copy
import json
from pathlib import Path

from docx import Document
from docx.image.image import Image
from docx.shared import Cm
from lxml import etree

from ..paths import FORMAT_SPEC, LIBRARY_DIR, TEMPLATE_DOCX
from . import omml
from .template_layout import clear_examples
from .model import Equation, Figure, FigureRow, Fragment, Heading, PageBreak, Para, Picture, Table

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
PIC = 'http://schemas.openxmlformats.org/drawingml/2006/picture'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'

PARA_STYLES = {'body': '报告-正文', 'plain': '报告-无缩进', 'center': '报告-无缩进', 'note': '报告-表注',
               'equation': '报告-公式', 'cover_note': '报告-封面注记'}
TOC_INSTRUCTION = ' TOC \\o "1-3" \\h \\z \\u '
TOC_PLACEHOLDER = '目录待更新：在 Word 中右键单击此处并选择“更新域”。'
TNR = 'Times New Roman'
MAX_FIGURE_CM = 16.0


def q(tag, ns=W):
    return '{%s}%s' % (ns, tag)


def format_spec():
    return json.loads(FORMAT_SPEC.read_text(encoding='utf-8'))


def _sub(parent, tag, **attrs):
    return etree.SubElement(parent, q(tag), {q(k): str(v) for k, v in attrs.items()})


class _Renderer:
    def __init__(self, document, spec):
        self.document = document
        # Word/WPS may rewrite style IDs on save; names are the upload contract.
        names = {style.name.lower(): style.style_id for style in document.styles}
        self.styles = {name: names.get(name.lower(), ident) for name, ident in spec['styles'].items()}
        self.table_spec = spec['table']
        body_section = spec['sections'][-1]
        self.text_width = body_section['pgSz']['w'] - body_section['pgMar']['left'] - body_section['pgMar']['right']
        # A 100% table extends by the cell margins on both sides (reference grid: 9070 + 2 × 108 = 9286 twips).
        self.table_width = self.text_width + 2 * self.table_spec['cell_margin_lr']

    # ---- primitives -------------------------------------------------------------------------
    def paragraph(self, style=None, jc=None):
        p = etree.Element(q('p'))
        ppr = _sub(p, 'pPr')
        if style:
            _sub(ppr, 'pStyle', val=self.styles[style])
        if jc:
            _sub(ppr, 'jc', val=jc)
        return p

    def text_run(self, text, bold=False, size=None, latin=None, italic=False, vert=None, color=None):
        r = etree.Element(q('r'))
        rpr = _sub(r, 'rPr')
        if latin:
            _sub(rpr, 'rFonts', ascii=latin, hAnsi=latin, cs=latin)
        if bold:
            _sub(rpr, 'b')
            _sub(rpr, 'bCs')
        if italic:
            _sub(rpr, 'i')
        if color:
            _sub(rpr, 'color', val=color)
        if size:
            _sub(rpr, 'sz', val=size)
            _sub(rpr, 'szCs', val=size)
        if vert:
            _sub(rpr, 'vertAlign', val=vert)
        if len(rpr) == 0:
            r.remove(rpr)
        t = _sub(r, 't')
        t.text = text
        t.set(XML_SPACE, 'preserve')
        return r

    def add_runs(self, p, runs, bold=False):
        for run in runs:
            if run.math is not None:
                p.append(copy.deepcopy(run.math))
                continue
            vert = 'subscript' if run.sub else 'superscript' if run.sup else None
            p.append(self.text_run(run.text, bold=bold or run.bold, italic=run.italic, vert=vert))
        return p

    def picture_run(self, path, width_cm):
        image = Image.from_file(str(path))
        native_cm = image.px_width / image.horz_dpi * 2.54
        width = min(width_cm or native_cm, MAX_FIGURE_CM)
        inline = self.document.part.new_pic_inline(str(path), Cm(width), None)
        r = etree.Element(q('r'))
        _sub(r, 'drawing').append(inline)
        return r

    # ---- cover and contents -----------------------------------------------------------------
    def blank_cover_line(self, size=None, jc=None):
        p = self.paragraph(jc=jc)
        rpr = _sub(p.find(q('pPr')), 'rPr')
        _sub(rpr, 'rFonts', ascii=TNR, hAnsi=TNR, cs=TNR)
        if size:
            _sub(rpr, 'sz', val=size)
            _sub(rpr, 'szCs', val=size)
        return p

    def cover(self, meta):
        cover = [self.blank_cover_line() for _ in range(8)]
        for line in meta['cover_lines']:
            p = self.paragraph('封面1')
            r = self.text_run(line, bold=True, size=52)
            r.find(q('rPr')).insert(0, etree.fromstring(f'<w:rFonts xmlns:w="{W}" w:cs="{TNR}"/>'))
            p.append(r)
            cover.append(p)
        cover += [self.blank_cover_line() for _ in range(11)]
        cover.append(self.blank_cover_line(size=30, jc='center'))
        for line in (meta['organization'], meta['date']):
            p = self.paragraph(jc='center')
            p.append(self.text_run(line, size=30, latin=TNR))
            cover.append(p)
        if meta.get('cover_note'):
            p = self.paragraph('报告-封面注记')
            p.append(self.text_run(meta['cover_note']))
            cover.append(p)
        return cover

    def toc(self):
        title = self.paragraph('报告-目录标题')
        title.append(self.text_run('目  录'))
        field = self.paragraph('toc 1')
        tabs = etree.SubElement(field.find(q('pPr')), q('tabs'))
        _sub(tabs, 'tab', val='right', leader='dot', pos=8504)
        begin = etree.SubElement(field, q('r'))
        _sub(begin, 'fldChar', fldCharType='begin')
        instr = _sub(etree.SubElement(field, q('r')), 'instrText')
        instr.text = TOC_INSTRUCTION
        instr.set(XML_SPACE, 'preserve')
        _sub(etree.SubElement(field, q('r')), 'fldChar', fldCharType='separate')
        field.append(self.text_run(TOC_PLACEHOLDER))
        _sub(etree.SubElement(field, q('r')), 'fldChar', fldCharType='end')
        return [title, field]

    # ---- blocks -----------------------------------------------------------------------------
    def render(self, block):
        if isinstance(block, Heading):
            return [self.heading(block)]
        if isinstance(block, Para):
            return [self.para(block)]
        if isinstance(block, Figure):
            return self.figure(block)
        if isinstance(block, FigureRow):
            return [self.figure_row(block)]
        if isinstance(block, Table):
            return self.table(block)
        if isinstance(block, Equation):
            return [self.equation(block)]
        if isinstance(block, Fragment):
            return self.fragment(block.name)
        if isinstance(block, Picture):
            p = self.paragraph('报告-公式')
            if block.image and Path(block.image).is_file():
                p.append(self.picture_run(block.image, block.width_cm))
            return [p]
        if isinstance(block, PageBreak):
            p = self.paragraph()
            _sub(etree.SubElement(p, q('r')), 'br', type='page')
            return [p]
        raise TypeError(f'未知报告块：{type(block).__name__}')

    def heading(self, block):
        p = self.paragraph(f'heading {block.level}')
        title = f'{block.label} {block.title}' if block.label else block.title
        p.append(self.text_run(title))
        return p

    def para(self, block):
        p = self.paragraph(PARA_STYLES[block.style], jc='center' if block.style == 'center' else None)
        return self.add_runs(p, block.runs)

    def caption(self, prefix, label, text, style):
        p = self.paragraph(style)
        p.append(self.text_run(f'{prefix}{label} {text}'))
        return p

    def figure_body(self, image, width_cm, caption):
        p = self.paragraph('报告-图')
        if image and Path(image).is_file():
            p.append(self.picture_run(image, width_cm))
            return p
        borders = _sub(p.find(q('pPr')), 'pBdr')
        for edge in ('top', 'left', 'bottom', 'right'):
            _sub(borders, edge, val='single', sz=4, space=4, color='808080')
        p.append(self.text_run(f'【图件待补充】{caption}', color='808080', size=21))
        return p

    def subcaption(self, text):
        p = self.paragraph('报告-资料表文')
        _sub(p.find(q('pPr')), 'keepNext')
        p.append(self.text_run(text))
        return p

    def figure(self, figure):
        out = []
        subcaptions = figure.subcaptions or []
        for n, (image, width) in enumerate(figure.images()):
            out.append(self.figure_body(image, width, figure.caption))
            if n < len(subcaptions) and subcaptions[n]:
                out.append(self.subcaption(subcaptions[n]))
        out.append(self.caption('图', figure.label, figure.caption, '报告-图题'))
        return out

    def figure_row(self, row):
        widths = [self.table_width // len(row.figures)] * len(row.figures)
        tbl = self.table_shell(widths, borders=False)
        tr = etree.SubElement(tbl, q('tr'))
        for figure, width in zip(row.figures, widths):
            tc = self.cell_shell(tr, width)
            for image, image_width in figure.images():
                tc.append(self.figure_body(image, image_width, figure.caption))
            tc.append(self.caption('图', figure.label, figure.caption, '报告-图题'))
        return tbl

    def table_shell(self, widths, borders=True):
        tbl = etree.Element(q('tbl'))
        pr = _sub(tbl, 'tblPr')
        _sub(pr, 'tblW', w=5000, type='pct')
        _sub(pr, 'jc', val='center')
        edges = _sub(pr, 'tblBorders')
        outer, inner = self.table_spec['outer'], self.table_spec['inner']
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            if not borders:
                _sub(edges, edge, val='nil')
            elif edge.startswith('inside'):
                _sub(edges, edge, val=inner['val'], sz=inner['sz'], space=0, color='auto')
            else:
                _sub(edges, edge, val=outer['val'], sz=outer['sz'], space=0, color='auto')
        _sub(pr, 'tblLayout', type='autofit')
        margins = _sub(pr, 'tblCellMar')
        _sub(margins, 'left', w=self.table_spec['cell_margin_lr'], type='dxa')
        _sub(margins, 'right', w=self.table_spec['cell_margin_lr'], type='dxa')
        _sub(pr, 'tblLook', val='04A0', firstRow=1, lastRow=0, firstColumn=1, lastColumn=0, noHBand=0, noVBand=1)
        grid = _sub(tbl, 'tblGrid')
        for width in widths:
            _sub(grid, 'gridCol', w=width)
        return tbl

    def cell_shell(self, tr, width, span=1, vmerge=None, nowrap=False):
        tc = etree.SubElement(tr, q('tc'))
        pr = _sub(tc, 'tcPr')
        _sub(pr, 'tcW', w=width, type='dxa')
        if span > 1:
            _sub(pr, 'gridSpan', val=span)
        if vmerge == 'restart':
            _sub(pr, 'vMerge', val='restart')
        elif vmerge == 'continue':
            _sub(pr, 'vMerge')
        if nowrap:
            _sub(pr, 'noWrap')
        _sub(pr, 'vAlign', val='center')
        return tc

    def table_widths(self, block):
        columns = sum(cell.span for cell in (block.header[0] if block.header else block.rows[0]))
        weights = block.widths or [1] * columns
        total = sum(weights)
        widths = [int(self.table_width * w / total) for w in weights]
        widths[-1] += self.table_width - sum(widths)
        return widths

    def table(self, block):
        out = []
        if block.caption:
            out.append(self.caption('表', block.label, block.caption, '报告-表题'))
        widths = self.table_widths(block)
        tbl = self.table_shell(widths)
        check = block.kind == 'check'
        head_style = '报告-表头' if check else '报告-资料表文'
        body_style = '报告-表文' if check else '报告-资料表文'
        head_height = self.table_spec['header_row_height'] if check else 360
        body_height = self.table_spec['body_row_height'] if check else 360
        for rows, style, height, is_head in ((block.header, head_style, head_height, True), (block.rows, body_style, body_height, False)):
            for row in rows:
                tr = etree.SubElement(tbl, q('tr'))
                trpr = _sub(tr, 'trPr')
                _sub(trpr, 'trHeight', val=height, hRule='atLeast')
                if is_head:
                    _sub(trpr, 'tblHeader')
                _sub(trpr, 'jc', val='center')
                col = 0
                for cell in row:
                    width = sum(widths[col:col + cell.span])
                    col += cell.span
                    tc = self.cell_shell(tr, width, cell.span, cell.vmerge, cell.nowrap)
                    tc.append(self.add_runs(self.paragraph(style), cell.runs, bold=cell.bold))
        out.append(tbl)
        out.extend(self.para(note) for note in block.notes)
        return out

    def equation(self, block):
        if not block.numbered:
            p = self.paragraph('报告-公式')
            p.append(omml.display(copy.deepcopy(block.omml)))
            return p
        number_width = self.table_width * 15 // 100
        widths = [self.table_width - number_width, number_width]
        tbl = self.table_shell(widths, borders=False)
        tr = etree.SubElement(tbl, q('tr'))
        formula = self.paragraph('报告-公式')
        formula.append(omml.display(copy.deepcopy(block.omml)))
        self.cell_shell(tr, widths[0]).append(formula)
        number = self.paragraph('报告-无缩进', jc='right')
        number.append(self.text_run(f'（{block.label}）'))
        self.cell_shell(tr, widths[1]).append(number)
        return tbl

    def fragment(self, name):
        root = etree.parse(str(LIBRARY_DIR / 'fragments' / f'{name}.xml')).getroot()
        out = []
        for child in root:
            element = copy.deepcopy(child)
            for blip in element.iter(q('blip', A)):
                ref = blip.get(q('embed', R))
                if ref and ref.startswith('LIB:'):
                    rid, _ = self.document.part.get_or_add_image(str(LIBRARY_DIR / 'figures' / ref[4:]))
                    blip.set(q('embed', R), rid)
            out.append(element)
        return out

    # ---- finishing ----------------------------------------------------------------------------
    def fill_placeholders(self, values):
        for part in self.document.part.package.iter_parts():
            name = str(part.partname)
            if not (name.startswith('/word/header') or name.startswith('/word/footer')):
                continue
            for t in part.element.iter(q('t')):
                for key, value in values.items():
                    if t.text and key in t.text:
                        t.text = t.text.replace(key, value)

    def renumber_drawings(self):
        body = self.document.element.body
        for n, drawing in enumerate(body.iter(q('drawing')), start=1):
            for tag in (q('docPr', WP), q('cNvPr', PIC)):
                for element in drawing.iter(tag):
                    element.set('id', str(n))


def write_docx(doc, path, template_path=None):
    """Write the finalized report to `path` using the lab master template."""
    spec = format_spec()
    document = Document(str(template_path or TEMPLATE_DOCX))
    renderer = _Renderer(document, spec)
    body = document.element.body
    cover_holder, toc_holder = clear_examples(document)
    final_sect = body.find(q('sectPr'))

    cover = renderer.cover(doc.meta)
    for element in cover:
        cover_holder.addprevious(element)
    cover_sect = cover_holder.find(q('pPr') + '/' + q('sectPr'))
    cover[-1].find(q('pPr')).append(cover_sect)
    body.remove(cover_holder)

    for element in renderer.toc():
        toc_holder.addprevious(element)
    for block in doc.blocks:
        for element in renderer.render(block):
            final_sect.addprevious(element)

    renderer.fill_placeholders({'{{HEADER}}': doc.meta['header_text'], '{{ORG}}': doc.meta['organization']})
    renderer.renumber_drawings()
    props = document.core_properties
    props.title = doc.meta.get('title', '')
    props.subject = doc.meta.get('subject', '')
    props.author = doc.meta['organization']
    props.last_modified_by = '工程报告撰写 Agent'
    props.keywords = doc.meta.get('keywords', '')
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return Path(path)


SETTINGS_AFTER_UPDATE_FIELDS = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids', 'mathPr',
                                'attachedSchema', 'themeFontLang', 'clrSchemeMapping', 'doNotIncludeSubdocsInStats',
                                'doNotAutoCompressPictures', 'forceUpgrade', 'captions', 'readModeInkLockDown',
                                'smartTagType', 'schemaLibrary', 'shapeDefaults', 'doNotEmbedSmartTags',
                                'decimalSymbol', 'listSeparator')


def set_update_fields(path):
    """Ask Word to refresh the table of contents on open (used when Word automation is unavailable)."""
    document = Document(str(path))
    settings = document.settings.element
    existing = settings.find(q('updateFields'))
    if existing is None:
        existing = etree.Element(q('updateFields'))
        anchor = next((c for c in settings if etree.QName(c).localname in SETTINGS_AFTER_UPDATE_FIELDS), None)
        if anchor is not None:
            anchor.addprevious(existing)
        else:
            settings.append(existing)
    existing.set(q('val'), 'true')
    document.save(str(path))
