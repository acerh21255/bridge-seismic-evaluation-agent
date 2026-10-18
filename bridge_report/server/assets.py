"""Local source inspection and SAP text geometry; never runs an uploaded program."""
import hashlib
import io
import json
import math
import re
import subprocess
import uuid
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document
from ..paths import TEMPLATE_DOCX, ROOT

PAIRS = re.compile(r'(\w+)=(?:"([^"]*)"|([^\s]+))')


def sap_geometry(text):
    tables, table = {}, None
    for line in text.replace(' _\n', ' ').splitlines():
        if line.startswith('TABLE:'):
            table = line.split('"')[1]
            tables.setdefault(table, [])
        elif table:
            row = {k: a or b for k, a, b in PAIRS.findall(line)}
            if row:
                tables[table].append(row)
    unit = next((r['CurrUnits'] for rows in tables.values() for r in rows if 'CurrUnits' in r), '')
    length = unit.split(',')[1].strip() if ',' in unit else ''
    scale = {'in': .0254, 'ft': .3048, 'mm': .001, 'cm': .01, 'm': 1}.get(length)
    if scale is None:
        raise ValueError('SAP 模型缺少可识别长度单位，不能生成有量纲三维视图')
    joints, warnings = {}, []
    for r in tables.get('JOINT COORDINATES', []):
        if r.get('CoordSys', 'GLOBAL') != 'GLOBAL' or r.get('CoordType', 'Cartesian') != 'Cartesian':
            warnings.append('非全局笛卡尔坐标节点未显示')
            continue
        xyz = [float(r.get(k, 0)) * scale for k in ('XorR', 'Y', 'Z')]
        if not all(math.isfinite(n) for n in xyz):
            raise ValueError('SAP 节点包含非有限坐标')
        joints[r['Joint']] = xyz
    edges = []
    for kind in ('FRAME', 'LINK'):
        for r in tables.get('CONNECTIVITY - ' + kind, []):
            if r.get('JointI') in joints and r.get('JointJ') in joints:
                edges.append({'id': r.get(kind.title()), 'a': r['JointI'], 'b': r['JointJ'], 'kind': kind.lower()})
            else:
                warnings.append('存在端点不完整的杆件')
    if not joints or not edges:
        raise ValueError('未找到 SAP 节点和杆件连接表，请导出 .s2k 或 .$2k 文本模型')
    if len(joints) > 50000 or len(edges) > 100000:
        raise ValueError('模型过大，请上传简化的可视化导出')
    if tables.get('CONNECTIVITY - AREA') or tables.get('CONNECTIVITY - SOLID'):
        warnings.append('当前视图仅展示杆件与连接单元，不含面/实体单元')
    return {'nodes': joints, 'edges': edges, 'units': 'm', 'source_units': unit,
            'warnings': sorted(set(warnings)), 'label': 'SAP 文本模型几何重建（未运行求解）'}


def template_validate(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024 or len(z.infolist()) > 1000:
                raise ValueError('模板解压体积或文件数超限')
            for name in z.namelist():
                if 'vbaproject' in name.lower() or '/embeddings/' in name:
                    raise ValueError('模板不得包含宏或嵌入对象')
                if name.endswith('.rels') and b'TargetMode="External"' in z.read(name):
                    raise ValueError('模板不得引用外部文件或链接')
        doc = Document(io.BytesIO(raw))
        base = Document(str(TEMPLATE_DOCX))
        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        body = doc.element.body
        from ..report.template_layout import section_holders
        section_holders(doc)
        required = [s.name for s in base.styles if s.name.startswith('报告-')]
        if any(s not in doc.styles for s in required):
            raise ValueError('模板缺少报告专用样式，请基于系统母版修改')
    except (zipfile.BadZipFile, KeyError) as exc:
        raise ValueError('不是可用的 DOCX 模板') from exc


class AssetService:
    def __init__(self, root):
        self.root = Path(root) / 'intake'
        self.root.mkdir(parents=True, exist_ok=True)

    def upload(self, raw, name):
        name = Path(name.replace('\\', '/')).name
        if not name or len(name) > 180:
            raise ValueError('文件名无效')
        ext = Path(name).suffix.lower()
        if ext not in ('.pdf', '.s2k', '.$2k', '.sdb', '.sap', '.xlsx', '.xls', '.docx', '.json', '.png', '.jpg', '.jpeg'):
            raise ValueError('支持图纸 PDF、SAP 文本模型、Excel、Word、JSON 和图像')
        ident = uuid.uuid4().hex
        folder = self.root / ident
        folder.mkdir()
        path = folder / name
        path.write_bytes(raw)
        record = {'id': ident, 'name': name, 'bytes': len(raw), 'sha256': hashlib.sha256(raw).hexdigest(),
                  'kind': '参考资料', 'status': '已归档', 'detail': '', 'preview': None}
        try:
            if ext in ('.s2k', '.$2k'):
                text = raw.decode('utf-8-sig', errors='replace')
                geometry = sap_geometry(text)
                (folder / 'geometry.json').write_text(json.dumps(geometry), encoding='utf-8')
                record.update(kind='有限元模型', status='已识别', geometry=f'/api/intake/{ident}/geometry',
                              detail=f"{len(geometry['nodes'])} 节点 · {len(geometry['edges'])} 杆件/连接单元 · 单位 m")
            elif ext in ('.sdb', '.sap'):
                record.update(kind='SAP 原生模型', status='待转换', detail='请在 SAP2000 导出 .s2k / .$2k；不把二进制文件当成已解析模型')
            elif ext == '.pdf':
                import pypdfium2 as pdfium
                with pdfium.PdfDocument(path) as pdf:
                    page = pdf[0]
                    scale = min(2, 2400 / max(page.get_size()))
                    page.render(scale=scale).to_pil().convert('RGB').save(folder / 'preview.jpg', quality=90)
                    texts = []
                    for n in range(min(len(pdf), 5)):
                        with pdf[n].get_textpage() as tp:
                            texts.append(tp.get_text_range())
                    text = '\n'.join(texts)
                    record.update(kind='图纸 / PDF', pages=len(pdf), preview=f'/api/intake/{ident}/preview',
                                  detail=f'{len(pdf)} 页；预览第 1 页，提取前 5 页文字')
                if len(text.strip()) < 30:
                    script = ROOT / 'tools' / 'drawing_ocr.ps1'
                    result = subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(script),
                                             '-Path', str(folder / 'preview.jpg')], capture_output=True, timeout=45,
                                            creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
                    text = result.stdout.decode('utf-8-sig', errors='replace') if result.returncode == 0 else ''
                    record['detail'] += '；首页 Windows OCR（需人工核对）' if text else '；扫描件文字识别不可用'
                record['extracted_text'] = text[:12000]
                record['candidates'] = {k: sorted(set(re.findall(pattern, text)))[:8] for k, pattern in {
                    '桩号': r'(?:\d{1,3})?K\s*\d{2,3}\s*\+\s*\d{3}',
                    '桥型': r'(?:预应力混凝土|钢筋混凝土|混凝土)?(?:连续梁桥|连续刚构桥|简支梁桥|空心板桥|T梁桥)',
                    '跨径候选': r'\d+(?:\.\d+)?\s*[+×x]\s*\d+(?:\.\d+)?(?:\s*\+\s*\d+(?:\.\d+)?)?'
                }.items()}
                record['status'] = '待核对'
            elif ext in ('.xlsx', '.xls'):
                record.update(kind='验算表格', detail='作为上游证据归档；数值须映射到交接合同后才能参与验算')
            elif ext == '.json':
                record.update(kind='交接数据', detail='请通过交接包入口加载并校验后使用')
        except Exception as exc:
            record.update(status='待补正', detail=str(exc)[:250] if isinstance(exc, ValueError) else '文件识别失败，请核对格式或导出方式')
        (folder / 'record.json').write_text(json.dumps(record, ensure_ascii=False), encoding='utf-8')
        return record

    def get(self, ident, kind):
        if not re.fullmatch('[a-f0-9]{32}', ident) or kind not in ('preview', 'geometry', 'record'):
            raise ValueError('资料不存在')
        path = self.root / ident / {'preview': 'preview.jpg', 'geometry': 'geometry.json', 'record': 'record.json'}[kind]
        return path

    def pdf_page(self, ident, number, recognize=False):
        import pypdfium2 as pdfium
        record_path = self.get(ident, 'record')
        record = json.loads(record_path.read_text(encoding='utf-8'))
        folder = record_path.parent
        if Path(record['name']).suffix.lower() != '.pdf':
            raise ValueError('该资料不是 PDF')
        if isinstance(number, bool) or not isinstance(number, int) or not 1 <= number <= record['pages']:
            raise ValueError('页码超出范围')
        preview = folder / f'page-{number}.jpg'
        if not preview.exists():
            with pdfium.PdfDocument(folder / record['name']) as pdf:
                page = pdf[number-1]
                page.render(scale=min(2.5, 3000/max(page.get_size()))).to_pil().convert('RGB').save(preview, quality=92)
        if not recognize:
            return preview
        with pdfium.PdfDocument(folder / record['name']) as pdf:
            with pdf[number-1].get_textpage() as tp:
                text = tp.get_text_range()
        method = 'PDF 文字层'
        if len(text.strip()) < 30:
            result = subprocess.run(['powershell.exe', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ROOT/'tools/drawing_ocr.ps1'),
                                     '-Path', str(preview)], capture_output=True, timeout=45,
                                    creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
            if result.returncode:
                raise ValueError('本页 OCR 未完成，请查看图纸并人工核对')
            text = result.stdout.decode('utf-8-sig', errors='replace'); method = 'Windows OCR'
        candidates = {k: sorted(set(re.findall(pattern, text)))[:12] for k, pattern in {
            '桩号': r'(?:\d{1,3})?K\s*\d{2,3}\s*\+\s*\d{3}',
            '桥型': r'(?:预应力混凝土|钢筋混凝土|混凝土)?(?:连续梁桥|连续刚构桥|简支梁桥|空心板|双柱墩|箱梁)',
            '跨径候选': r'\d+(?:\.\d+)?\s*[+×x]\s*\d+(?:\.\d+)?(?:\s*\+\s*\d+(?:\.\d+)?)?'
        }.items()}
        result = {'name':record['name'],'id':ident,'page':number,'method':method,'extracted_text':text[:18000],
                  'candidates':candidates,'message':'识别结果仅作资料索引，桥型与尺寸须结合原图核对；未写入验算数据。'}
        (folder/f'page-{number}-ocr.json').write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
        return result
