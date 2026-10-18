# -*- coding: utf-8 -*-
"""
筛选数据表 -> 桥梁参数 JSON 批量转换器
适配 G8511筛选结果.xlsx 的表格结构（Subagent2 输出格式）：
  - 每座桥占 2 行（首行=基本信息+左柱墩高，次行=右柱墩高）
  - "墩高" 标签列之后：均值列、空列、逐墩墩高值
用法:
  python xlsx_to_params.py 输入表.xlsx(或.csv) --outdir ./params [--ag 0.10] [--site II]
输出: 每桥一个 <桩号>.json（BridgeParams 兼容）+ _conversion_summary.json
"""
import argparse
import csv
import json
import math
import os
import re

# 主梁等效截面预设（待图纸核定，仅供建模演示）
DECK_PRESETS = {
    "空心板": {"width": 12.25, "area": 7.5, "inertia_33": 2.5, "inertia_22": 25.0,
               "torsion_j": 4.0, "shear_as2": 6.3, "shear_as3": 6.3},
    "T梁":    {"width": 12.0, "area": 6.0, "inertia_33": 4.0, "inertia_22": 30.0,
               "torsion_j": 5.5, "shear_as2": 5.0, "shear_as3": 5.0},
    "箱梁":   {"width": 12.75, "area": 8.5, "inertia_33": 6.0, "inertia_22": 60.0,
               "torsion_j": 12.0, "shear_as2": 7.0, "shear_as3": 7.0},
}
MATERIALS = {
    "C50": {"name": "C50", "fck": 32.4, "elastic_mod": 34500},
    "C30": {"name": "C30", "fck": 20.1, "elastic_mod": 30000},
    "C25": {"name": "C25", "fck": 16.7, "elastic_mod": 28000},
}
SKIP_TYPES = ["拱", "斜拉", "悬索"]  # v1 不支持桥型


def read_table(path):
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig") as f:
            return [r for r in csv.reader(f)]
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    return [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]


def parse_spans(combo: str):
    """'4*30' / '1*20+4*30' / '11+24+11' -> [float,...]"""
    combo = combo.replace("×", "*").replace("x", "*").replace("X", "*").strip()
    spans = []
    for part in combo.split("+"):
        part = part.strip()
        m = re.match(r"^(\d+)\s*\*\s*(\d+\.?\d*)$", part)
        if m:
            spans += [float(m.group(2))] * int(m.group(1))
        else:
            m2 = re.match(r"^(\d+\.?\d*)$", part)
            if m2:
                spans.append(float(m2.group(1)))
    return spans


def num(s, default=None):
    m = re.search(r"\d+\.?\d*", str(s))
    return float(m.group()) if m else default


def parse_pier_size(s):
    """'1.5' -> ('circle',1.5)；'5.6*5.6'/'2.6*1.5' -> ('rect',等效直径)"""
    s = str(s).replace("×", "*")
    if "*" in s:
        a, b = s.split("*")[:2]
        a, b = num(a), num(b)
        if a and b:
            return "rect", round(math.sqrt(4 * a * b / math.pi), 3), (a, b)
    d = num(s)
    return ("circle", d, None) if d else (None, None, None)


def convert(table_path, outdir, ag=0.10, intensity=7, site="II", tg=0.45):
    rows = read_table(table_path)
    os.makedirs(outdir, exist_ok=True)
    bridges, warnings_all = [], []
    i = 1
    while i < len(rows):
        r = rows[i]
        is_start = bool(str(r[1]).strip() or str(r[3]).strip())
        if not is_start:
            i += 1
            continue
        r2 = rows[i + 1] if i + 1 < len(rows) else []
        # 次行可能是"右柱墩高行"（序号/桩号为空），也可能是下一座桥的首行（本桥无次行）
        r2_is_start = bool(r2) and bool(str(r2[1]).strip() or str(r2[3]).strip())
        warn = []

        # 在数据行内精确定位 "墩高" 标签（表头里"纵向墩高差"等会误命中）
        gh_col = next((k for k, v in enumerate(r) if str(v).strip() == "墩高"), None)
        if gh_col is None:
            i += 1 if r2_is_start else 2
            continue
        h_start = gh_col + 3  # 标签列 + 均值列 + 空列 -> 逐墩墩高起始列

        seq = str(r[1]).strip()
        stake = str(r[3] or r[2]).strip().replace("/", "_")
        file_key = f"{seq}_{stake}" if seq else stake
        side = str(r[4]).strip()
        spans = parse_spans(str(r[6]))
        girder = str(r[8]).strip()
        structure = str(r[9]).strip()
        bearings = str(r[10]).strip()
        abut_type = str(r[11]).strip() or "重力式桥台"
        pier_type = str(r[12]).strip()
        pile_type = str(r[16]).strip()
        pile_d = num(r[17], 1.5)
        cap_w = num(r[18], 1.7)

        if any(t in structure + pier_type for t in SKIP_TYPES):
            warnings_all.append(f"{stake}: 桥型({structure}) v1 不支持，跳过")
            i += 1 if r2_is_start else 2
            continue
        if not spans or len(spans) < 2:
            warnings_all.append(f"{stake}: 跨径组合解析失败({r[6]})，跳过")
            i += 1 if r2_is_start else 2
            continue

        shape, d_eq, rect = parse_pier_size(r[13])
        if not d_eq:
            warnings_all.append(f"{stake}: 桥墩尺寸缺失，跳过")
            i += 1 if r2_is_start else 2
            continue
        if shape == "rect":
            warn.append(f"矩形墩 {rect[0]}×{rect[1]}m 按等效圆截面 φ{d_eq}m 近似")

        n_cols = 1 if "单柱" in pier_type else 2
        # 逐墩墩高：首行与次行取均值
        def heights(row):
            out = []
            for c in range(h_start, min(len(row), h_start + len(spans) + 2)):
                v = num(row[c])
                if v is not None and v > 0.5:
                    out.append(v)
            return out
        h1 = heights(r)
        h2 = heights(r2) if (r2 and not r2_is_start) else []
        n_piers = len(spans) - 1
        pier_hs = []
        for k in range(n_piers):
            vals = [h for h in (h1[k] if k < len(h1) else None,
                              h2[k] if k < len(h2) else None) if h]
            pier_hs.append(round(sum(vals) / len(vals), 2) if vals
                             else round(num(r[gh_col + 1], 10.0), 2))
        if len(h1) < n_piers:
            warn.append(f"墩高数据不足({len(h1)}/{n_piers})，缺失墩用均值补齐")

        deck_key = next((k for k in DECK_PRESETS if k in girder), "T梁")
        deck = dict(DECK_PRESETS[deck_key], girder_type=girder,
                    unit_weight=26.0, material="C50")
        L = sum(spans)
        x = 0.0
        piers = []
        for k in range(n_piers):
            x += spans[k]
            piers.append({
                "pier_id": f"P{k+1}", "location": round(x, 2), "height": pier_hs[k],
                "n_columns": n_cols, "column_diameter": d_eq, "column_spacing": 6.0,
                "cap_beam_height": 1.5, "cap_beam_width": cap_w,
                "pile_diameter": pile_d, "pile_length": 25.0, "soil_m": 10000.0,
                "column_material": "C30", "pile_material": "C25"})

        data = {
            "bridge_name": f"{stake} 桥（{side}）" if side else f"{stake} 桥",
            "stake": stake,
            "source_drawing": f"{os.path.basename(table_path)} 序号{r[1]}",
            "spans": [{"length": s} for s in spans],
            "piers": piers,
            "abutments": [{"location": 0.0, "type": abut_type, "restraint": "纵向活动"},
                          {"location": L, "type": abut_type, "restraint": "纵向活动"}],
            "deck": deck,
            "materials": MATERIALS,
            "seismic": {"intensity": intensity, "ag": ag, "site_class": site,
                        "tg": tg, "ci": 1.0, "cs": 1.0, "cd": 1.0, "damping": 0.05},
            "bearings": bearings,
            "pile_type": pile_type,
            "_note": "柱间距6.0m/桩长25m/主梁等效截面/设防参数为默认假定，需图纸与地勘核定",
            "_warnings": warn,
        }
        out = os.path.join(outdir, f"{file_key}.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        bridges.append({"stake": stake, "seq": seq, "spans": spans, "n_piers": n_piers,
                        "pier_heights": pier_hs, "warnings": warn, "json": out})
        i += 1 if r2_is_start else 2  # 有次行吃两行，无次行只过一行

    summary = {"converted": len(bridges), "bridges": bridges,
               "skipped_or_warnings": warnings_all}
    with open(os.path.join(outdir, "_conversion_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print(f"转换完成: {len(bridges)} 座桥 -> {outdir}")
    for w in warnings_all:
        print("  [跳过]", w)
    return summary


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("table")
    ap.add_argument("--outdir", default="./params")
    ap.add_argument("--ag", type=float, default=0.10)
    ap.add_argument("--intensity", type=int, default=7)
    ap.add_argument("--site", default="II")
    ap.add_argument("--tg", type=float, default=0.45)
    args = ap.parse_args()
    convert(args.table, args.outdir, args.ag, args.intensity, args.site, args.tg)
