import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from bridge_report.paths import FORMAT_SPEC, LIBRARY_DIR
from bridge_report.report import omml
from bridge_report.report.docx_writer import set_update_fields, write_docx
from bridge_report.report.model import (Cell, Equation, Figure, FigureRow, Fragment, Heading, Para, ReportDoc, Run, Table,
                                        finalize)

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
STYLES = json.loads(FORMAT_SPEC.read_text(encoding='utf-8'))['styles']


def q(tag, ns=W):
    return '{%s}%s' % (ns, tag)


def text(el):
    return ''.join(t.text or '' for t in el.iter(q('t'), q('t', M)))


def sample_doc():
    meta = {'cover_lines': ['G8511昆磨高速（抗震韧性提升工程）', '设计桩号 K278+308', '20m +33m + 20 m 跨径', '混凝土连续梁桥', '抗震验算分析报告'],
            'organization': '土木工程防灾减灾全国重点实验室', 'date': '2026.03', 'cover_note': '注：演示变体。',
            'header_text': 'G8511昆磨高速设计桩号 K278+308  20+33+20 m连续梁桥', 'title': '抗震验算分析报告'}
    check = Table('t', 'E1地震作用下桥墩验算结果(纵桥向)',
                  header=[[Cell([Run('地震作用')]), Cell([Run('关键截面')]), Cell([Run('验算结果')])]],
                  rows=[[Cell([Run('纵向地震')], vmerge='restart'), Cell([Run('1#墩底')]), Cell([Run('通过')])],
                        [Cell([], vmerge='continue'), Cell([Run('2#墩底')]), Cell([Run('强度破坏')], bold=True)]],
                  notes=[Para([Run('注：示例表注。')], style='note')], widths=[1, 1, 1])
    blocks = [Heading(1, '工程概况'), Heading(2, '桥梁概况'),
              Para([Run('截面承载力见'), Run(ref='t'), Run('，符号'), Run('M'), Run('y', sub=True), Run('。')]),
              Figure('f', '等效屈服弯矩示意', image=LIBRARY_DIR / 'figures' / 'm_phi.png', width_cm=8.3),
              Figure('missing', '缺失的立面图'),
              check,
              FigureRow([Figure('a', '桩基截面', image=LIBRARY_DIR / 'figures' / 'm_phi.png', width_cm=6),
                         Figure('b', '桥墩截面', image=LIBRARY_DIR / 'figures' / 'm_phi.png', width_cm=6)]),
              Equation(omml.math('V≤', omml.sub('V', 'cs')), key='e', numbered=True),
              Equation(omml.math('a=b')),
              Heading(4, '（1）纵向输入', numbered=False),
              Fragment('shear_formulas')]
    return finalize(ReportDoc(meta=meta, blocks=blocks))


class DocxWriterTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = Path(cls.tmp.name) / 'report.docx'
        write_docx(sample_doc(), cls.path)
        cls.zip = zipfile.ZipFile(cls.path)
        cls.body = etree.fromstring(cls.zip.read('word/document.xml')).find(q('body'))
        cls.paras = list(cls.body.iter(q('p')))

    @classmethod
    def tearDownClass(cls):
        cls.zip.close()
        cls.tmp.cleanup()

    def style_of(self, p):
        ps = p.find(q('pPr') + '/' + q('pStyle'))
        return ps.get(q('val')) if ps is not None else None

    def test_opens_and_has_three_sections(self):
        self.assertEqual(len(Document(str(self.path)).sections), 3)
        self.assertEqual(len(list(self.body.iter(q('sectPr')))), 3)

    def test_template_examples_do_not_leak_into_generated_report(self):
        content = text(self.body)
        for example in ('目录与使用说明', '桥梁抗震验算报告母版', '{{构件名称}}', '{{项目名称}}', '正文样式示例'):
            self.assertNotIn(example, content)

    def test_visible_upload_preserves_edited_styles(self):
        from bridge_report.paths import TEMPLATE_DOCX
        from bridge_report.server.assets import template_validate
        from docx.shared import Pt
        with tempfile.TemporaryDirectory() as folder:
            master = Document(TEMPLATE_DOCX)
            master.styles['报告-正文'].font.size = Pt(13)
            master.styles['报告-正文'].element.set(q('styleId'), 'BRBodyEdited')
            template = Path(folder) / 'custom.docx'
            master.save(template)
            template_validate(template.read_bytes())
            output = Path(folder) / 'result.docx'
            write_docx(sample_doc(), output, template_path=template)
            result = Document(output)
            self.assertEqual(result.styles['报告-正文'].font.size.pt, 13)
            self.assertTrue(any(p.style.style_id == 'BRBodyEdited' for p in result.paragraphs))
            self.assertNotIn('正文样式示例', '\n'.join(p.text for p in result.paragraphs))

    def test_cover(self):
        cover = [text(p) for p in self.paras if self.style_of(p) == STYLES['封面1']]
        self.assertEqual(cover, sample_doc().meta['cover_lines'])
        first_sect_para = next(p for p in self.body.findall(q('p')) if p.find(q('pPr') + '/' + q('sectPr')) is not None)
        self.assertEqual(text(first_sect_para), '注：演示变体。')
        self.assertIn('2026.03', [text(p) for p in self.paras])

    def test_toc_field(self):
        instr = ''.join(t.text for t in self.body.iter(q('instrText')))
        self.assertIn('TOC \\o "1-3" \\h \\z \\u', instr)
        self.assertIn('目  录', [text(p) for p in self.paras])

    def test_headers_and_footers_filled(self):
        parts = {n: text(etree.fromstring(self.zip.read(n))) for n in self.zip.namelist() if n.startswith(('word/header', 'word/footer'))}
        joined = ''.join(parts.values())
        self.assertIn('G8511昆磨高速设计桩号 K278+308  20+33+20 m连续梁桥', joined)
        self.assertIn('土木工程防灾减灾全国重点实验室', joined)
        self.assertNotIn('{{', joined)

    def test_heading_text_and_styles(self):
        h1 = next(p for p in self.paras if self.style_of(p) == STYLES['heading 1'])
        h4 = next(p for p in self.paras if self.style_of(p) == STYLES['heading 4'])
        self.assertEqual((text(h1), text(h4)), ('1. 工程概况', '（1）纵向输入'))

    def test_captions_and_reference(self):
        texts = [text(p) for p in self.paras]
        self.assertIn('图1.1-1 等效屈服弯矩示意', texts)
        self.assertIn('表1.1-1 E1地震作用下桥墩验算结果(纵桥向)', texts)
        self.assertIn('截面承载力见表1.1-1，符号My。', texts)
        self.assertIn('图1.1-3 桩基截面', texts)
        self.assertIn('图1.1-4 桥墩截面', texts)

    def test_missing_figure_placeholder(self):
        texts = [text(p) for p in self.paras]
        self.assertTrue(any('待补充' in t and '缺失的立面图' in t for t in texts))
        self.assertIn('图1.1-2 缺失的立面图', texts)

    def test_check_table_format(self):
        table = next(t for t in self.body.iter(q('tbl')) if '强度破坏' in text(t))
        borders = table.find(q('tblPr') + '/' + q('tblBorders'))
        self.assertEqual((borders.find(q('top')).get(q('val')), borders.find(q('top')).get(q('sz'))), ('double', '6'))
        self.assertEqual((borders.find(q('insideH')).get(q('val')), borders.find(q('insideH')).get(q('sz'))), ('single', '4'))
        self.assertEqual(table.find(q('tblPr') + '/' + q('tblW')).get(q('w')), '5000')
        merges = [m.get(q('val')) for m in table.iter(q('vMerge'))]
        self.assertEqual(merges, ['restart', None])
        bold_cell = next(tc for tc in table.iter(q('tc')) if text(tc) == '强度破坏')
        self.assertIsNotNone(next(bold_cell.iter(q('b')), None))
        header = table.find(q('tr'))
        self.assertIsNotNone(header.find(q('trPr') + '/' + q('tblHeader')))

    def test_numbered_equation(self):
        tables = [t for t in self.body.iter(q('tbl')) if '（1-1）' in text(t)]
        self.assertEqual(len(tables), 1)
        self.assertIsNotNone(next(tables[0].iter(q('oMath', M)), None))

    def test_fragment_inserted(self):
        self.assertTrue(any('偏心受压构件的斜截面受剪承载力' in text(p) for p in self.paras))

    def test_images_embedded(self):
        media = [n for n in self.zip.namelist() if n.startswith('word/media/')]
        self.assertTrue(media)
        self.assertEqual(len(list(self.body.iter('{http://schemas.openxmlformats.org/drawingml/2006/picture}pic'))), 3)

    def test_update_fields_flag(self):
        path = Path(self.tmp.name) / 'flag.docx'
        write_docx(sample_doc(), path)
        set_update_fields(path)
        with zipfile.ZipFile(path) as z:
            settings = etree.fromstring(z.read('word/settings.xml'))
        self.assertEqual(settings.find(q('updateFields')).get(q('val')), 'true')


if __name__ == '__main__':
    unittest.main()
